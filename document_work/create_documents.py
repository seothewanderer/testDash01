"""Three review documents derived from the shipped dashboard and original plan."""
from pathlib import Path
from copy import deepcopy
import csv, json, hashlib, zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'document_work'
OUT=ROOT/'deliverables'/'문서정리_20260927'
OUT.mkdir(exist_ok=True)
REF=Path('C:/Users/user/Downloads/드론_진로탐색_대시보드_프로젝트_기획서.docx')
CAT={r['name']:r for r in json.loads((ROOT/'dashboard/data/catalog.json').read_text(encoding='utf8'))}
EV=json.loads((WORK/'runtime-evidence.json').read_text(encoding='utf8'))
S=EV['sources']
AUD=json.loads((ROOT/'dashboard/data/audit.json').read_text(encoding='utf8'))
DATE='2026년 9월 27일'

# Each entry records actual fields and implementation behavior, not planned joins.
SPEC={
'I':('산업 규모 연도별 추이','연도별 산업 집계 4행','reference_year','2021~2024년',
 'reference_year는 조사 기준연도, company_total은 업체 수, employees_total은 종사자 수, revenue_total_100m_krw는 억원 단위 매출이다. valid_sample_total, comparability_flag, quality_note는 비교 범위와 주석이다.',
 '3개 추이 차트와 2024년 지표를 표시한다. 2024년 업체 6,493개, 세부표 기준 매출 11,077.19억원, 종사자 17,204명이다. 개인 필터는 적용하지 않는다.',
 '산업의 구조와 변화를 파악한 다음 활용 분야와 직무를 탐색하게 한다. 조종 외 개발·제조·활용 업무로 관심을 넓히는 효과를 기대한다.',
 '연도별 표본·정의 변화와 매출 요약문 불일치가 있다. 종사자 수를 채용 인원으로 바꾸거나 증감 원인을 경기로 단정하지 않는다.'),
'U':('공공기관 활용 사례','자산 그룹별 활용 기록 30행','asset_group_id','as_of_date 필드 기준',
 'purpose_raw는 활용 목적 원문, purpose_category는 목적 분류이다. quantity는 보유 수량, purchase_year는 구매연도이며 as_of_date와 source_file은 기준일과 출처다.',
 '산업 화면의 실제 공공 활용 사례에서 purpose_raw 중복 제거 후 앞 6개를 보여준다. 수량을 합산한 전국 공공부문 보유량 차트는 만들지 않았다.',
 '연구 분야 막대를 실제 활용 업무의 문장으로 보완하고 직무 탐색으로 연결한다. 추상적인 분야명이 실제 업무로 어떻게 나타나는지 이해하도록 돕는다.',
 '원자료의 기관·자산 범위를 넘어 일반화하지 않는다. 구매연도와 자료 기준일은 현재 운용 상태나 채용 시점이 아니다.'),
'R':('국가 연구개발 과제','고유 연구과제 8,078개','project_id','대표 시작연도와 최신 기준연도는 과제별 상이',
 'ranking_inclusion_flag는 기존 고신뢰 포함 여부, defense_flag 등 활용 flag와 ai_autonomous_flight_flag 등 기술 flag는 다중 태그이다. project_title, lead_institution, representative_start_year, research_objective_summary는 사례 설명에 쓰인다.',
 '기본 고신뢰 2,936개에서 분야별 과제 수를 집계한다. 국방 태그는 818개다. 선택 분야와 연구 범위에 맞춰 6개 기술 막대를 보여주고 제목에 드론·무인·UAV가 있는 앞 5개 사례를 펼친다.',
 '관심 활용 분야 → 연구 문제와 기술 → 국방 기술 또는 직무 탐색으로 연결한다. 배워 볼 기술과 프로젝트의 업무 맥락을 제공한다.',
 '기존 분류를 외부 재검증한 결과가 아니다. 다중 태그 합계는 고유 과제 수를 초과할 수 있다. 과제 수·연구비는 채용 수요나 취업 확률이 아니다.'),
'D':('국방 계약 사례','계약 항목 416행과 고유 계약번호 367개','record_id 및 contract_no','contract_date 필드 기준',
 'contract_no는 중복 제거 단위, contract_agency는 계약 기관, contract_title은 업무명, contractor는 계약자다. contract_date와 contract_status는 계약 맥락을 설명한다.',
 '기관별 고유 contract_no 수의 상위 7개를 막대로 표시하고 계약 사례를 펼친다. 동일 번호가 기관 간 나타날 수 있으므로 기관별 수를 전체 계약 수로 단순 합산하지 않는다.',
 '국방 사업이 개발뿐 아니라 납품·정비·용역 업무로 나타남을 설명한다. 국방 기술과 기업 근거를 함께 탐색할 업무 맥락을 제공한다.',
 '계약 금액을 일자리 수로 변환하지 않는다. 계약자와 기업 마스터의 이름을 새로 자동 연결하지 않았으며 기업의 현재 채용을 뜻하지 않는다.'),
'DT':('국방 드론 기술 분류','기술 분류 12개','technology_id','원본 source_status 및 출처 기준',
 'technology_category는 기술 분류, keywords는 기술 단어, defense_use_case는 국방 활용 설명이다. source_status, source_file, source_page, quality_note는 출처와 해석 조건이다.',
 '국방 기술 선택 메뉴와 응용 설명을 제공한다. 동일 technology_id의 TM을 연결하여 역할과 기초 학습 내용을 표시한다.',
 '방산 관심 → 기술의 활용 맥락 → 관련 역할과 학습으로 이어지게 한다. 방산이라는 명칭보다 구체적으로 준비할 내용을 판단하도록 돕는다.',
 '기술 분류는 실제 채용 빈도나 공식 자격 요건이 아니다. 기술의 존재만으로 특정 기업의 채용 조건을 추정하지 않는다.'),
'TM':('국방 기술과 직무 학습 매핑','기술별 분석용 매핑 12행','technology_id','mapping_type 및 usage_note 기준',
 'job_group과 example_roles는 관련 역할 설명, job_keywords는 직무 후보 검색어, learning_topics는 기초 학습 주제다. mapping_type과 usage_note는 매핑의 성격을 설명한다.',
 'DT와 직접 연결한다. 관련 직무 보기에서는 job_keywords와 learning_topics의 명시 단어로 T의 제목·업무·기술을 검색한다. 역할 규칙의 tech 목록으로 준비 화면에 학습 내용을 보여준다.',
 '국방 응용 기술에서 사전 직무와 공통 기초학습으로 연결한다. 전공·기술을 어떤 업무에 활용할 수 있는지 검토할 수 있다.',
 '분석용 매핑과 문자열 후보 검색이다. 검증된 추천 순위나 국방 공고의 요구 역량 통계로 해석하지 않는다.'),
'T':('드론 직무 분류 사전','직무 역할 209개','job_id','collected_date 및 근거 필드 기준',
 'major_category는 8개 원본 대분류, job_title_ko는 직무명, core_duties는 업무, skills_raw는 기술 원문이다. preferred_qualifications, workplace_types_raw, employer_examples_raw는 자격·근무처·기업 예시이며 evidence_type은 근거 성격이다.',
 '분류 타일과 활동·분야·검색 필터로 후보를 좁힌다. 선택 직무의 업무, 근무처, 역량, 직접 공고 사례를 표시한다. TS·TE는 job_id로 연결하고 J·E 후보는 별도 문자열 규칙으로 연결한다.',
 '산업 관심 → 역할 선택 → 준비 역량 → 공고 조건의 중심 기준으로 사용한다. 조종 외 개발·설계·정비·사업 역할을 탐색하도록 돕는다.',
 '209개는 현재 채용 직업 수가 아니다. 직접 근거 31개·보조 근거 20개·파생 직무 158개가 섞여 있다. 사전의 기업 예시도 현재 채용 기업을 뜻하지 않는다.'),
'TS':('직무별 준비 역량 관계','직무와 기술 관계 645행','job_id와 skill_order 및 skill_keyword','T에 종속된 사전 기준',
 'job_id는 T 연결키, skill_keyword는 사전 기술명, skill_order는 순서다. 계산 시 직무별 skill_keyword의 중복을 제거한 집합을 사용한다.',
 '보유 기술 체크와 미체크 목록을 만들고 정확히 같은 기술명 교집합이 큰 주변 직무 최대 4개를 비교한다. 선택 직무 포함 최대 5행·기술 최대 9열의 명시 여부 행렬을 표시한다.',
 '직무 선택 → 보유 기술과 추가 학습 → 교육 후보·프로젝트·로드맵으로 연결한다. 이미 가진 기반 지식과 보완할 기술을 구분하도록 돕는다.',
 '기술 체크는 자기보고이며 숙련도나 합격률 평가가 아니다. 동의어 정규화가 없어 같은 의미의 다른 표기가 일치하지 않을 수 있다. 교육 커리큘럼과 직접 연결된 자료도 아니다.'),
'TE':('직무와 연결된 채용 사례','직무별 직접 연결 사례 20행','job_id 및 actual_posting_url','사례 원문 시점과 현재 상태 별개',
 'job_id는 사전 직무, actual_employer는 기업명, actual_posting_title과 actual_posting_url은 공고 제목과 원문이다. salary_original_text 등 급여 원문도 보유한다.',
 '선택 직무의 상세 더보기에서 기업·공고 제목·원문 링크를 제공한다. J의 공고 수에 20행을 더하지 않으며 J와 공고 동일성을 새로 확정하지 않는다.',
 '직무 설명에서 실제 관측된 채용 업무 사례로 이동하는 직접 근거를 제공한다. 문자열 후보 매핑과 직접 연결 사례를 구분할 수 있다.',
 '작은 표본의 사례다. 직접 연결이 없는 직무를 채용이 없는 직무로 판단하지 않는다. 현재 모집 여부는 확인되지 않았다.'),
'J':('수집 채용 공고 분석','고유 공고 135개와 기업명 76개','posting_id','게시일과 마감일 및 수집일 미보유',
 'job_major_category는 채용 대분류, province_name은 근무지역, career_type과 career_min_years/max는 경력 조건이다. education_normalized/raw, responsibilities, major_requirement_raw, salary_raw, primary_source_url은 조건·업무·출처다.',
 '지역·직무·경력·학력 분포와 공고 목록, 방산 근거 흐름도, 개인 조건 비교에 사용한다. 기업명 정확 일치로 C에 52개를 연결하고 미연결 83개를 유지한다. 원문 연봉 9개만 다시 파싱한다.',
 '학습 후보 → 관측된 기업·공고 조건 → 로드맵의 확인 행동으로 연결한다. 추상적인 준비를 구체적인 조건과 미확인 항목으로 바꾸도록 돕는다.',
 '채용 시장 전체 또는 현재 모집 목록이 아니다. 업무 원문 103건·전공 11건으로 불완전하다. 미기재를 무관으로 바꾸지 않으며 현재 지원 가능 여부와 합격률은 계산하지 않는다.'),
'K':('전체 공고 키워드 빈도','키워드 집계 189행과 언급 합 622','keyword_category_name과 keyword_normalized','수집 공고 135개 전체 표본',
 'keyword_category_name은 범주, keyword_normalized는 표준화 표현, posting_count는 언급 공고 수, posting_share_pct는 전체 표본 비율이다. posting_id 연결키는 없다.',
 '6개 기술·도구 범주에서 빈도 상위 12개를 막대로 보여준다. 자격증 범주의 원본 앞 6개도 별도 표시한다. 직무·지역·방산 필터는 적용하지 않는다.',
 '사전 준비 역량 옆에 전체 채용 언어를 참고자료로 제공한다. 사전의 기술과 관측 공고의 언급을 구분해 검토할 수 있다.',
 '622는 공고 수가 아니라 다중 키워드 언급의 합이다. J의 keyword_count 합과 같아도 공고별 연결은 복원되지 않는다. 직무별 기술 빈도나 동시 등장 조합을 만들지 않는다.'),
'E':('교육과정 마스터','고유 과정 148개','course_id','회차 날짜는 ES 참조',
 'course_name, ncs_name, related_certificate_raw는 후보 검색 근거다. institution_name, total_training_hours, remote_mode는 기관·시간·방식이다. tuition_won, actual_training_cost_won, government_subsidy_won, expected_out_of_pocket_won은 비용 원문이다.',
 '직무 규칙의 단어가 과정명·NCS·관련 자격에 있는 경우 후보로 표시한다. 산점도 가로축은 총 훈련시간이다. 과정 선택을 로드맵에 반영하고 비용 네 값을 원문 이름 그대로 펼친다.',
 '직무 준비 → 명시 근거가 있는 과정 후보 → 커리큘럼·비용 확인 → 준비 계획으로 연결한다. 실제로 알아볼 교육 기회를 좁힐 수 있다.',
 '147/148과정의 비용 산술관계가 불일치한다. 지원액을 차감해 부담액을 확정하지 않는다. 이름이 일치해도 특정 기술을 실제로 배우는지와 취업 성과는 미확인이다.'),
'ES':('교육 회차 일정','개설 회차 1,262개','offering_id 및 course_id','시작일은 2026년이며 마지막 종료일은 2027년 6월 17일',
 'course_id는 E 연결키, offering_id는 회차키, province_name과 remote_mode는 지역·방식이다. start_date/end_date는 일정, course_url은 상세 원문이다. 만족도·수료율 필드는 이번 화면 집계에 사용하지 않는다.',
 '선택 지역 또는 원격 회차를 유지한다. 과정별 산점도 세로축은 현재 범위 회차 수, 월별 막대는 2026년 개강 회차 수다. 일정 상태는 2026년 9월 26일 기준 종료 916·진행 129·예정 217이다.',
 '과정 후보 → 실제 일정·지역 → 개인 준비 계획으로 연결한다. 과정 수와 개설 기회를 나누어 현실적인 학습 일정을 검토할 수 있다.',
 '과정 148개와 회차 1,262개는 다른 단위다. 일정상 예정은 현재 모집 중이라는 뜻이 아니다. 신청률·만족도·수료율을 취업률로 대체하지 않는다.'),
'Q':('연간 자격 취득 통계','자격 종류와 연도별 집계 80행','certificate_type과 year 및 period_label','차트는 완결 연도 중 마지막 6개인 2017~2022년',
 'certificate_type은 자격 종류, year는 연도, acquisition_count는 취득 건수다. full_year는 완결 연도 여부, period_end는 기간 끝, measure_type와 source_unit은 지표 정의다.',
 '무인멀티콥터 중 full_year=Y 자료의 최근 6개 연도 취득 추이를 그린다. K의 자격 언급 막대와 나란히 놓고 두 모집단을 직접 연결하지 않는다.',
 '자격 취득 규모와 수집 공고의 언급을 구분해 자격·교육을 탐색하게 한다. 많이 취득되는 자격을 곧바로 취업 효과가 큰 자격으로 오해하지 않도록 돕는다.',
 '고유 인원·취업자 수·활동 조종사 수·합격률이 아니다. 최신 현황이나 법적 필수 요건 및 개인 응시 자격을 판정하지 않는다.'),
'C':('기업과 기관 마스터','고유 기업과 기관 300개','company_id 및 company_name_normalized','원본 분류 근거와 CD의 시점 기준',
 'company_name_normalized는 J와의 정확 일치 비교 이름이다. defense_drone_relation_class는 원문 관계 분류, classification_basis는 이유다. official_defense_company_designation_verified와 designation_note는 공식 지정 검증 상태다.',
 'J의 employer_name과 정확히 같은 이름만 연결한다. 기업 카드·방산 집단별 공고 흐름도·전체 300개 기업의 분류 구성에 사용한다. 미분류 201개와 공고 미연결 83개를 서로 구분한다.',
 '공고 기업의 사업 성격 → 방산 근거 → 관측된 직무와 조건으로 연결한다. 기업의 방산 연관성과 개인 지원 조건을 분리해서 판단할 수 있다.',
 'N 300개는 공식 지정 미검증이며 지정업체가 아니라는 뜻이 아니다. 미분류를 민간으로 바꾸지 않는다. 별칭이나 법인 표기를 자동 정규화하여 연결 범위를 늘리지 않았다.'),
'CB':('기업별 사업 분야 관계','중복 정리 사업 관계 1,322행','record_id 및 company_id','source_dataset과 source_row 기준',
 'company_id는 C 연결키, business_category는 사업 분야, defense_relation은 분야와 국방의 관계다. source_dataset/source_row/record_json은 원본 근거를 담는다.',
 '기업 카드에 사업 분야를 표시한다. 선택 공고에 정확히 연결된 기업의 company_id와 business_category 쌍을 중복 제거해 상위 6개 분야의 기업 수를 집계한다.',
 '채용 기업이 어떤 일을 하는지 확인하고 직무와의 사업 연관성을 검토하게 한다. 채용 조건만으로 부족한 기업 이해를 보완한다.',
 '관계행 수는 기업 수가 아니다. 한 기업이 여러 분야에 나타날 수 있어 막대 합계를 전체 기업 수로 보지 않는다. 분야명만으로 공식 방산 지정을 추정하지 않는다.'),
'CD':('기업별 방산 근거','근거 기록 106행','record_id 및 evidence_id와 company_id','record_json 내부 as_of_date 등 원문 기준',
 'evidence_type은 근거 종류, evidence_strength는 강도, official_source는 출처 성격, caveat는 주의점이다. record_json의 evidence_detail/as_of_date/source_file을 함께 읽는다.',
 'C와 company_id로 연결해 기업 더보기에서 근거 종류·강도·시점·주의점을 보여준다. 공고의 방산 집단 분류 자체는 C의 원본 분류를 사용한다.',
 '방산 관련 집단 → 구체적인 기업 근거 → 개별 공고 확인으로 연결한다. 관계가 확인된 범위를 검토하고 미확인 사항을 남길 수 있다.',
 '파일명 Verified가 최신 원문 전수 검증을 뜻하지 않는다. 근거 106행을 기업 수나 채용 수로 합산하지 않는다. 국적·보안심사·군 경력 조건을 기업 근거에서 추론하지 않는다.'),
}

def font(run,size=None,bold=None):
    run.font.name='맑은 고딕'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    run.font.color.rgb=RGBColor(0,0,0)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold

def new_doc(title,intro):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(1.8)
    for name,size in [('Normal',10.5),('Title',23),('Heading 1',15),('Heading 2',12),('Heading 3',10.5)]:
        st=d.styles[name];st.font.name='맑은 고딕';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=Pt(size*1.4)
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.widow_control=True
    d.add_paragraph(title,'Title');p(d,f'{DATE}  |  시각화 개선판 기준',size=9)
    p(d,intro)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=foot.add_run();font(run,9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    d.core_properties.author='드론 진로 탐색 대시보드 프로젝트'
    return d

def p(d,text,bold=False,size=None):
    x=d.add_paragraph();r=x.add_run(text);font(r,size,bold);return x
def h(d,text,level=1):return d.add_heading(text,level)
def page(d,title):
    heading=h(d,title)
    heading.paragraph_format.page_break_before=True
def label(d,name,text):
    x=d.add_paragraph();font(x.add_run(name+'  '),bold=True);font(x.add_run(text));return x
def table(d,headers,data,widths=None,size=9.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for row in data:
        for c,x in zip(t.add_row().cells,row):c.text=str(x)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDEAF2' if ri==0 else ('F7F9FB' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                a=OxmlElement('w:'+edge);a.set(qn('w:w'),'75');a.set(qn('w:type'),'dxa');mar.append(a)
            cp.append(mar)
            bor=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                a=OxmlElement('w:'+edge);a.set(qn('w:val'),'single');a.set(qn('w:sz'),'4');a.set(qn('w:color'),'D9D9D9');bor.append(a)
            cp.append(bor)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=Pt(size*1.35)
                for run in para.runs:font(run,size,ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

GAPS=[
('특허와 일반 조달 등 산업 보조 맥락','기획안 5절 산업 단계의 특허·조달, 설계 A04~A06·4절 산업 상세','특허 순위, PPS 일반 조달, 전망 O, 산업 원시 IR, 수출 X, 예산 상세는 미채택이다. 연구 R와 국방 계약 D 중심으로 구현했다.','현재 선별 범위로 명시한다. 필요한 보조 질문을 정한 뒤 추가 검토하며 원래 기획 항목을 삭제하지 않는다.'),
('관심 분야의 수량 의미','설계 A02의 활용 매출·공공 응답과 사례','현재 분야 막대는 R 고신뢰 연구의 활용 태그 수이고 U의 활용 목적은 펼침 사례다. 산업 활용 시장 규모 차트는 아니다.','연구 태그로 보는 분야 탐색이라는 범위를 유지한다. 매출·공공 응답을 원하면 IR 정제가 필요하다.'),
('진입 조건과 개인 조건 비교','설계 D02·D08·J02·J04 및 5번째 화면','경력·학력 구성과 선택 공고의 학력·경력·지역 비교는 구현했다. 직무별 신입/경력 2축 산점도, 최소 경력 분포, 사용자 전공·고용형태 입력 비교는 미구현이다.','단순화된 테스트 범위를 명시한다. 입력과 원문 해석을 검토한 후 확장하며 완전한 개인 매칭으로 표현하지 않는다.'),
('직무와 과정 및 공고 연결','기획안 6절 직무 중심 연결, 설계 7절 다대다 연결표','T 제목/분류에 처음 맞는 ROLE_RULES 규칙으로 J의 넓은 대분류와 E 검색어를 정한다. 검토된 독립 다대다 매핑 원장은 없다.','후보 연결로 유지한다. 역할별 다중 매핑·예외·근거를 검토해 원장화하는 작업은 후속이다.'),
('직접 기술 문구와 방산 특수 조건','설계 E06·E07·I05의 제한적 원문 사례','공고 업무·전공·경력 원문은 열람 가능하다. 공고별 기술 및 병역 문구를 구조화해 재추출하지 않았고 기술 빈도·특수 조건 영역은 빈 결과로 두었다.','완전한 관계표 부재와 원문 일부 보유를 구분한다. 추가 추출과 검토가 필요한 상태로 기록한다.'),
('여러 목표와 세부 교육 비교','설계의 관심 직무 1~2개, G06·G07, 지역 시나리오','현재 직무·과정·공고는 각각 하나를 선택한다. 교육 만족도·수료율·난도 태그 및 지역 교육/공고 통합 비교는 별도 구현하지 않았다.','목표 다중 저장과 보조 비교를 후속 후보로 남긴다. 원본 교육 평가지표 보유를 데이터 부재로 쓰지 않는다.'),
]

def unused_reason(name):
    if name in ['Company_Examples_By_Job_Long','Workplace_Type_By_Job_Long']:return 'T의 원문 예시·근무처 필드로 우선 설명'
    if name.startswith('NTIS_Prior'):return '이전 자료 대조 이력이며 독립 표본 아님'
    if name.startswith('NTIS_'):return 'R의 과제·태그 중심 사용; 상세 관계 확장 보류'
    if name in ['Company_Keyword_Matrix','Company_Keyword_Relations_Long','Company_Count_By_Keyword']:return 'C와 기업 ID 체계가 달라 별도 매핑 필요'
    if 'Patent' in name:return '독립 순위 보조 맥락으로 보류; 기업 연결 미확정'
    if name=='PPS_Drone_Bids':return '일반 조달 보조 맥락 보류; D 국방 계약 우선'
    if name in ['Employment_Outlook','Export_Competitiveness_Trade','Industry_Survey_Raw_Data']:return '산업 보조 맥락 보류; 단위·정의별 정제 필요'
    if 'Budget' in name:return '예산 보조 맥락 보류; 채용으로 환산하지 않음'
    if name in ['2024_Cumulative_Certifications','Pilot_Certifications_By_Aircraft','Pilot_Certifications_By_Period','Private_Certification_Status']:return '연간 Q 계열 우선; 누적·등록 자료와 중복 방지'
    if name in ['Public_Sector_Drone_Inventory','Registered_Operator_Activities','Supply_Demand_By_Region','Swarm_Drone_Registered_Companies']:return '등록·지역 보조 맥락 보류; 관측 단위가 다름'
    return '채택 마스터·검증본 또는 원문 필드와 중복·보조 관계'

def build_spec():
    d=new_doc('드론 진로 탐색 대시보드\n데이터 명세와 활용 설계','제작자와 팀원이 현재 화면의 숫자와 연결 근거를 검토하기 위한 문서다. 원본 CSV 54개 중 핵심 진로 흐름에 필요한 17개를 화면에 채택했으며, 나머지 37개는 원본을 보존했다. 기대효과는 설계상 기대이며 사용자 조사로 측정된 성과가 아니다.')
    h(d,'1 채택 데이터와 관측 단위')
    p(d,'데이터 기준일은 2026년 9월 26일이다. 개별 통계·공고·과정·기업 근거의 원래 시점은 서로 다르다. 아래 코드는 현재 구현 코드의 SOURCES를 따른다. 설계보고서의 E·DT 등 묶음 코드와 달리 과정/회차 및 기술/매핑을 분리했다.')
    table(d,['코드','데이터','원본 규모'],[(k,v[0],f"{CAT[S[k]]['rows']:,}행") for k,v in SPEC.items()],[1.2,11.5,4.7])
    p(d,'명세 범위: 화면 계산·표시·연결에 사용하는 핵심 필드와 제한을 기록했다. 미사용 열을 포함한 전체 컬럼 목록·결측 수·SHA-256은 프로젝트의 dashboard/data/catalog.json에 보존되어 있다. 파일별 실제 이름은 다음 명세에 기록했다.',size=9)
    page(d,'2 데이터 연결과 해석 규칙')
    table(d,['연결','현재 방법','검토할 의미'],[
      ('T → TS·TE','job_id 직접 연결','사전 역량과 관측 공고 사례를 별도 표시'),('DT → TM','technology_id 직접 연결','분석용 기술·학습 매핑'),('E → ES','course_id 직접 연결','과정 수와 회차 수 분리'),('C → CB·CD','company_id 직접 연결','기업·사업 관계·근거의 단위 분리'),('J → C','employer_name = company_name_normalized','52공고 연결, 83공고 미연결 보존'),('T → J','직무명/분류의 ROLE_RULES 첫 일치 규칙','넓은 채용 분류 후보이며 직접 채용 확정 아님'),('T → E','규칙 단어와 과정명·NCS·자격의 포함 일치','TS 개별 기술과 교육 모듈의 직접 조인 아님'),('J → K','직접 연결 없음','전체 키워드 빈도를 개별 공고에 배분 금지')],[2.5,6.4,8.5])
    h(d,'집계와 필터',2)
    p(d,'관심 분야는 직무 후보와 연구 범위에 적용한다. 희망지역은 공고와 교육 회차에 적용하며 원격 과정은 유지한다. 지역 지도만 현재 지역 필터를 잠시 해제하고 나머지 공고 필터는 유지한다. 산업 전체 추이·전체 키워드·전체 기업 분류는 개인 조건으로 바뀌지 않는다.')
    p(d,'직무별 기술 집합은 정확한 문자열로 중복 제거한다. 연봉 파생값은 연봉으로 명시된 9개만 원문에서 다시 읽으며 두 범위 오류는 5,000~8,000만원과 2,600~5,000만원으로 처리한다. 원본 CSV 자체는 수정하지 않았다.')
    p(d,'관측 0건, 연결 미확인, 데이터 미수집은 서로 다르다. 공고와 기업, 과정과 회차, 기업과 사업 관계, 과제와 태그를 혼합 합산하지 않는다. 비용 불일치 147과정은 원본 금액을 나란히 보여 주며 확정 부담액을 계산하지 않는다.')
    entries=list(SPEC.items())
    for i in range(0,len(entries),2):
        page(d,f'3 핵심 데이터 명세 {i//2+1}')
        for code,v in entries[i:i+2]:
            title,unit,key,time,fields,use,purpose,limit=v
            h(d,code+' '+title,2)
            p(d,S[code]+'.csv',size=9)
            p(d,'관측 단위: '+unit+' · 식별: '+key+'\n자료 시점: '+time,size=9)
            label(d,'핵심 필드',fields)
            label(d,'현재 활용',use)
            label(d,'목적과 기대효과',purpose)
            label(d,'주의와 제한',limit)
            p(d,'원본 '+str(len(CAT[S[code]]['columns']))+'개 컬럼 · CSV 문자열을 읽고 집계 대상 값만 수치로 변환',size=8.5)
    page(d,'4 화면에서 데이터가 이어지는 방식')
    for title,keys in [
      ('산업 이해',['industry-size','applications','research','defense-tech','defense-contracts']),
      ('직무 탐색',['role-explore','role-detail','adjacent-roles']),
      ('준비 역량',['skills','keyword-reference','courses','portfolio','certifications']),
      ('수집 채용 공고',['posting-conditions','postings','companies','defense-comparison','defense-explore']),
      ('준비 로드맵',['action-plan','personal-input','condition-match'])]:
        h(d,title,2)
        for key in keys:
            m=EV['modules'][key]
            p(d,f"{m['title']}  [{', '.join(m['sources'])}] — {m['purpose']}.",size=9.5)
    p(d,'위 코드는 모듈 더보기의 관련 데이터 범위다. 예를 들어 교육 후보는 T 기반 규칙과 E·ES를 직접 계산하고 TS는 준비 역량 맥락으로 연결된다. 방산 공고 흐름도의 직접 계산은 J·C이며 CD는 분류 근거 열람에 쓰인다.',size=9)
    for idx,keys in enumerate([['live','keywords','curriculum'],['interview','defense','outcomes','mobility']]):
        page(d,f'5 추가 수집 데이터와 빈 결과 영역 {idx+1}')
        p(d,'다음은 구현된 빈 결과 영역이다. 기존 원문 일부 보유와 완성된 구조화 데이터셋의 부재를 구분한다. 수치·모의 그래프·성과값은 채우지 않았다.')
        for key in keys:
            m=EV['missing'][key];h(d,m['title'],2)
            label(d,'필요 항목',m['fields'])
            label(d,'연결과 활용',m['flow']+' / '+m['method'])
            label(d,'목적과 기대효과',m['purpose']+'. '+m['effect']+'.')
            p(d,m['priority']+' · '+m['collection'],size=9)
    unused=[(n,r) for n,r in CAT.items() if not r['selected']]
    for idx in range(0,len(unused),19):
        page(d,f'6 미채택 원본과 보류 이유 {idx//19+1}')
        p(d,'이 목록은 데이터가 없다는 뜻이 아니다. 현재 화면에 내장하지 않은 자료이며 원본 54개 전체는 프로젝트 ZIP의 data/raw에 보존되어 있다. 사후에 정리한 보류 이유로, 당시 개별 파일별 의사결정 시각을 재현하는 기록은 아니다.',size=9.5)
        table(d,['원본 파일명','행 수','현재 보류 이유'],[(n+'.csv',f"{r['rows']:,}",unused_reason(n)) for n,r in unused[idx:idx+19]],[8.0,1.4,8.0],8.5)
    page(d,'7 검토 위치와 근거 자료')
    h(d,'데이터를 선별하는 방법',2)
    p(d,'각 모듈의 더보기에서 사용 데이터·연결 흐름·목적·기대효과·계산 규칙·한계를 읽고 미검토/유지/수정/제외를 기록한다. 제외는 의견으로 저장되며 실제 화면을 자동으로 삭제하지 않는다. 원본 파일 버튼은 채택 데이터의 검색·페이지 이동·결측 정보를 제공한다.')
    p(d,'32개 검토 모듈에는 데이터 활용 모듈 21개, 추가 수집 빈 영역 7개, 제작자 검토 모듈 4개가 포함된다. 검토 메모와 선택 상태를 포함한 HTML 저장 또는 의견 JSON 저장이 가능하다. 자동 저장·공동 편집 서버는 없다.')
    table(d,['근거','확인할 내용'],[
      ('첨부 datas.zip','54개 CSV 원본 출처'),('첨부 대시보드 설계보고서','질문별 목적·관측 범위·한계'),('첨부 프로젝트 기획서','5단계 목표와 기대효과'),('dashboard/build.py','채택 17개·파생 급여·일정·품질 점검'),('dashboard/src/app.js','직무 후보·기업 연결·조건 비교·빈 데이터 정의'),('dashboard/src/charts.js 및 views.js','시각화 집계·필터·표현'),('dashboard/src/evidence.json','모듈 목적과 기대효과의 기본 메타데이터'),('dashboard/data/catalog.json 및 audit.json','행 수·컬럼·해시·연결 원장·품질 수치')],[7.0,10.4])
    p(d,'외부 원문 최신성, 기업 동일성 전체, 공식 방산 지정, 교육 품질을 이번 구현에서 새로 전수 검증하지 않았다. 문서의 현재 구현 설명은 2026년 9월 27일 생성 결과물을 기준으로 한다.')
    f=OUT/'01_데이터_명세와_활용설계.docx';d.save(f);return f

def build_report():
    d=new_doc('드론 진로 탐색 대시보드\n제작 과정과 기획 정합성 보고서','현재 결과물은 원래의 진로 탐색 목적과 5단계 흐름을 유지한 제작자 검토용 테스트 대시보드다. 핵심 데이터 17개를 연결했고, 시각화 개선을 통해 차트 선택이 다음 탐색으로 이어지도록 구성했다. 설계보고서 전체의 모든 상세 기능을 완성한 상태는 아니며, 아래에 차이와 남은 범위를 명시한다.')
    h(d,'1 작업 범위와 산출물')
    table(d,['입력','작업에서의 역할'],[
      ('프로젝트 기획서 DOCX','목적·대상·기대효과·5단계 흐름의 기준'),('대시보드 설계보고서 MD','질문별 데이터 연결·집계 제한·시나리오의 기준'),('datas.zip','54개 CSV 분석과 핵심 17개 선별'),('이전 제작 HTML','시각적 분포 탐색과 차트 구성의 참고')],[6.2,11.2])
    p(d,'사용자는 모든 데이터를 사용하는 것보다 기획에 적합한 데이터를 선별하고, 부족한 데이터의 활용 위치를 빈 영역으로 검토할 수 있도록 요청했다. 모든 모듈의 더보기와 방산 연관성, 가상환경·직접 테스트, HTML 공유는 필수 조건으로 반영했다.')
    h(d,'현재 산출물',2)
    p(d,'독립 실행 HTML, 재생성 소스와 원본 CSV를 포함한 프로젝트 ZIP, 사용 안내, 테스트 기록을 만들었다. 대시보드는 5단계 화면과 제작자 검토 화면으로 구성되며 검토 모듈은 32개다. 계정·서버 기반 저장·현재 공고 자동 수집은 구현 범위에 포함하지 않았다.')
    h(d,'기록의 기준',2)
    p(d,'이 보고서는 남아 있는 코드·데이터 원장·테스트 결과·대화 요청을 대조하여 작업의 논리적 순서와 결정 이유를 정리한 것이다. 단계별 소요시간이나 팀원의 참여, 사용자 효과 측정값은 기록이 없으므로 제시하지 않는다. 데이터 기준일은 2026년 9월 26일, 시각화 개선판 검증과 문서 정리는 9월 27일이다.')
    page(d,'2 데이터 검토부터 초기 구현까지')
    for title,text in [
      ('요구사항과 지침 확인','기획의 산업 이해 → 직무 탐색 → 준비 역량 → 실제 채용 → 취업 준비 로드맵 흐름을 기준으로 삼았다. AGENTS.md의 가정 명시·최소 구현·국소 수정·검증 원칙을 적용했고 dashboard/AGENTS.md에 데이터 단위, 방산 해석, 빈 영역과 HTML 저장 규칙을 기록했다.'),
      ('자료 목록과 품질 점검','54개 CSV를 파싱하여 행 수·컬럼·결측·해시를 기록했다. 공고 135개, 직무 209개, 교육 148과정과 1,262회차, 기업 300개의 단위를 구분했다. 기업 정확 일치 52공고와 미연결 83공고, 비용 불일치 147과정, 연봉 범위 오류 2건 등을 확인했다.'),
      ('필요한 데이터 선별','핵심 탐색을 연결하는 17개를 채택했다. 기업·역량·회차는 직접 키로 연결하고, 직무와 공고·교육처럼 직접 키가 없는 부분은 명시적인 후보 규칙을 사용했다. 특허·일반 조달·수출·산업 원시 세부 등은 원본을 보존하고 현재 화면의 범위를 좁혔다.'),
      ('검토 가능한 화면 구성','모듈별 사용 데이터·목적·기대효과·계산·한계와 검토 의견을 더보기에 배치했다. 최신 모집 상태·공고별 기술·커리큘럼 등 7개 영역에는 빈 결과와 필요한 필드·수집 경로를 만들었다. 기업 미확인을 민간으로, 공고 미기재를 무관으로 바꾸지 않았다.'),
      ('실행 환경과 공유 구조','프로젝트 상위 .venv에 Python 환경을 만들고 표준 라이브러리로 데이터를 파싱·압축·HTML에 내장했다. 화면은 HTML/CSS/JavaScript와 SVG로 구현했다. 외부 CDN 없이 HTML 한 파일로 열리며 선택과 메모를 포함한 복사본을 저장하도록 했다.')]:
        h(d,title,2);p(d,text)
    page(d,'3 시각화 개선과 사용자 흐름')
    p(d,'초기 결과물이 보고서 포털처럼 텍스트 중심이라는 피드백을 반영했다. 첨부 예시 HTML에서 지역 분포·구성비·흐름도·교육 분포 같은 탐색 방식을 참고하되, 현재 데이터가 뒷받침하지 않는 공고별 기술 연결·이분법 방산 비교·비용 합산은 사용하지 않았다.')
    table(d,['화면','현재 시각화','다음 행동'],[
      ('산업 이해','업체·매출·종사자 3개 추이, 분야·기술 막대, 국방 계약','관심 분야 또는 국방 기술에서 직무로 이동'),('직무 탐색','분류 타일, 역할 카드, 공통 기술 행렬','직무를 선택하고 역량과 공고 후보 확인'),('준비 역량','보유 기술 표시, 교육 산점도·월별 개강, 자격 추이','과정·프로젝트 후보를 로드맵에 선택'),('수집 채용 공고','지역 타일, 경력·학력·직무 분포, 방산 근거 흐름도','차트를 필터로 사용해 원문 공고 선택'),('준비 로드맵','목표·역량·학습·결과물·공고의 경로, 조건 상태','학습·과정·공고 확인 행동을 체크')],[3.0,7.4,7.0])
    h(d,'선택 상태의 흐름',2)
    p(d,'직무·활용 분야·희망지역은 화면 사이에서 유지한다. 직무를 변경하면 종속 과정·공고·프로젝트·완료 체크를 비우고 보유 기술과 검토 메모는 유지한다. 직무가 비어 있는 상태에서 경력 필터가 남던 초기화 오류를 시각화 테스트에서 발견해 수정했다.')
    h(d,'방산을 유지한 방법',2)
    p(d,'국방 기술 12개와 학습 매핑, 고신뢰 연구의 국방 태그, 국방 계약 업무, 기업 근거 종류·강도·시점을 연결했다. 공고 흐름도는 정확히 연결된 기업의 원본 방산 분류와 공고 직무를 사용한다. 미연결과 미분류도 별도 집단으로 표시하며 국가 예산이나 계약을 개인의 채용 가능성으로 환산하지 않는다.')
    page(d,'4 검증 결과와 남은 한계')
    table(d,['검증','확인한 결과'],[
      ('Python 데이터 테스트','7개 통과. 기본키·연결키, 관측 단위, 기업 연결 보존, 비용·일정, 연봉 원문 처리'),('브라우저 통합 검증','6개 화면, 32개 더보기, 직무→기술→과정→공고→로드맵, 원문 검색, 선택 변경과 7개 빈 영역'),('시각화 검증','지역 합계 135·경기 49, 방산 흐름선 합 135, 월별 회차 합, 직무 분류 필터, 교육 점의 키보드 선택'),('공유와 표시','HTML 저장·오프라인 복원, 모바일 390px·데스크톱·글자 200%에서 문서 전체 가로 넘침 없음'),('실행과 원본','Edge에서 JavaScript 오류 0건, 외부 HTTP 요청 0건. 원본 CSV 보존 및 패키지 HTML 동일성 확인')],[4.2,13.2])
    p(d,'실행 근거는 dashboard/tests/test_data.py, browser.cjs, visuals.cjs와 TEST_REPORT.md다. 스크립트는 각각 데이터 규칙, 주요 사용자 경로, 시각적 집계와 상호작용을 검증한다. 렌더링 이미지는 내부 검토용이며 이번 문서 폴더의 납품 파일에는 포함하지 않는다.')
    h(d,'검증되지 않은 범위',2)
    p(d,'직무별 후보 규칙의 의미 적합성을 전수 평가하지 않았다. 제목의 첫 일치 규칙은 복합 직무를 단순화할 수 있고, 정확한 기술명 비교는 동의어를 놓칠 수 있다. 외부 공고의 현재 상태, 기업 동일성 전체, 공식 방산 지정, 교육 비용 정의 및 교육 품질은 새로 검증하지 않았다. Firefox·Safari의 호환성도 별도 검증하지 않았다.')
    h(d,'기대효과와 실제 성과',2)
    p(d,'탐색 부담 완화와 준비 방향 구체화는 기획의 기대효과다. 이번 작업에서는 화면 동작과 집계의 일관성을 확인했으며, 이용자의 탐색 시간 감소·진로 선택 개선·취업 성과를 실험으로 확인한 것은 아니다. 후속 사용자 검토에서는 단계별 이탈·추천 이유 이해·다음 행동 결정 여부를 확인할 수 있다.')
    for i in range(0,len(GAPS),3):
        page(d,f'5 기획과 현재 구현의 차이 {i//3+1}')
        if i==0:p(d,'기획안의 1~7절은 목표와 서비스 방향을 규정한다. 현재 구현은 그 핵심을 유지하지만 상세 설계와 완전히 같지는 않다. 자료 선별에 대한 사용자 지시로 범위를 좁힌 항목과, 아직 구현·구조화하지 않은 기능을 구분해 아래에 남긴다. 미완료를 숨기기 위해 기획 목표를 낮추지 않는다.')
        for title,source,current,nextstep in GAPS[i:i+3]:
            h(d,title,2);label(d,'기준',source);label(d,'현재 차이',current);label(d,'처리',nextstep)
    page(d,'6 기획안 보완 내역과 후속 작업')
    p(d,'원래 기획안의 목적·대상·기대효과·5단계 흐름·데이터 원칙을 대체하거나 삭제할 필요는 없다. 보완본은 원문 1~7절을 그대로 두고 현재 구현 상태를 설명하는 8~11절을 뒤에 추가했다. 앞의 “실제 채용”은 현재 화면의 “수집 채용 공고”를 의미하며, 모집 중임을 뜻하지 않는다는 설명을 추가했다.')
    table(d,['추가 위치','추가한 내용'],[
      ('8 현재 테스트 구현 범위','17개 채택, 현재 화면과 선택·저장 범위, 데이터 시점'),('9 기획과 구현의 차이','특허·조달 미채택, 비교 기능 단순화, 규칙 후보 연결 및 미구현 구분'),('10 데이터 보완과 검증','7개 빈 영역, 우선 수집 3종, 완료한 테스트와 미검증 범위'),('11 보완 이력','원문 변경 0·삭제 0·추가 4절과 추가 사유')],[6.3,11.1])
    h(d,'후속 순서 제안',2)
    p(d,'먼저 제작자가 모듈의 유지·수정·제외와 직무 후보 규칙을 검토한다. 이어서 공고별 키워드 관계표·교육 커리큘럼·최신 모집 상태를 보완하면 현재 단계 사이의 연결을 가장 직접적으로 강화할 수 있다. 이후 특허·일반 조달·산업 지역 상세 및 비교 기능은 실제 사용자 질문에 필요한 항목부터 추가한다.')
    p(d,'이번 문서 작업은 대시보드 코드를 변경하거나 추가 데이터를 수집한 작업이 아니다. 현재 결과물의 사용 데이터·제작 과정·계획과 구현의 차이를 기록하고 원문 기획에 구현 현황을 보완했다.')
    f=OUT/'02_제작과정과_기획정합성_보고서.docx';d.save(f);return f

def build_plan():
    # Append only to the existing document body. Preserve every other ZIP part.
    d=Document(REF);body=d.element.body;original=list(body);original_nonsect=[e for e in original if e.tag!=qn('w:sectPr')]
    page(d,'8 현재 테스트 구현 범위')
    p(d,f'보완 기준일 {DATE}. 원문 1~7절의 목적·대상·기대효과·흐름은 유지한다. 이 절부터는 시각화 개선판의 현재 상태와 후속 검토 범위를 구분한 구현 보완 내용이다.')
    p(d,'현재 단계는 제작자가 데이터 활용을 선별하고 흐름을 검토하는 테스트 구현이다. 원본 CSV 54개 중 17개를 사용하고 37개는 보존했다. 모든 자료를 사용하는 것이 목표는 아니며, 각 모듈의 목적과 다음 단계 연결이 채택 기준이다.')
    table(d,['원래 단계','현재 구현','선택 결과'],[
      ('산업 이해','3개 산업 추이, 연구 분야·기술, 국방 기술·계약','관심 분야와 관련 직무 탐색'),('직무 탐색','분류 타일·활동 필터·직무 설명·공통 역량 행렬','관심 직무 1개'),('준비 역량','기술 체크, 교육 산점도·일정, 자격과 프로젝트','과정 1개와 프로젝트 제안'),('실제 채용','수집 채용 공고의 분포·기업 근거·방산 흐름도','비교 공고 1개'),('취업 준비 로드맵','학습 경로·조건 상태·다음 행동 3개','선택·보유 기술·확인 사항')],[3.2,8.2,6.3],9.5)
    p(d,'화면의 “수집 채용 공고”는 원문 기획의 “실제 채용”을 관측 자료의 범위에 맞게 표현한 명칭이다. 현재 모집 여부가 미확인인 공고를 포함하며 지원 가능 확정을 의미하지 않는다.')
    p(d,'시각화는 차트 선택과 다음 탐색에 집중한다. 사용 데이터·흐름·목적·기대효과·계산 규칙·한계·검토 의견은 모든 모듈의 더보기에 둔다. 근거 전체 펼치기, 원문 데이터 검색, 미검토/유지/수정/제외 판단과 메모를 제공한다.')
    p(d,'HTML 한 파일에 데이터와 코드를 내장하여 서버 없이 열람한다. “검토 내용 포함 HTML 저장”으로 현재 선택·개인 입력·보유 기술·메모를 팀원에게 전달할 수 있다. 자동 저장과 공동 편집 서버는 없다. 기준일은 2026년 9월 26일이며 원본 통계의 연도와 개별 자료의 시점은 따로 표시한다.')
    for i in range(0,len(GAPS),3):
        page(d,f'9 기획과 현재 구현의 차이 {i//3+1}')
        if i==0:p(d,'원문 단계별 데이터는 기획의 활용 후보 범위이며 모두 현재 구현되었다는 뜻은 아니다. 상세 설계의 미구현 사항을 완료로 바꾸거나 삭제하지 않고 다음과 같이 기록한다.')
        for title,source,current,nextstep in GAPS[i:i+3]:
            h(d,title,2);label(d,'기준',source);label(d,'현재 상태',current);label(d,'후속 처리',nextstep)
    page(d,'10 데이터 보완과 검증')
    p(d,'자료가 부족한 부분에는 빈 결과와 수집 제안을 배치한다. 최신 모집 상태·공고별 기술 관계·교육 커리큘럼을 우선 보완한다. 현직자 업무, 방산 특수 지원 조건, 교육·자격 성과, 실제 경력 이동은 후속 수집 영역으로 둔다. 기존 원문이 일부 존재하는 것과 구조화된 전체 관계표가 없는 것은 구분한다.')
    table(d,['추가 데이터','연결 목적과 기대효과'],[(EV['missing'][k]['title'],EV['missing'][k]['purpose']+'. '+EV['missing'][k]['effect']) for k in ['live','keywords','curriculum','interview','defense','outcomes','mobility']],[6.1,11.6],9.5)
    p(d,'집계와 직접 연결키, 연봉 원문·비용·일정의 데이터 테스트 7개를 통과했다. Edge에서 주요 탐색 경로, 32개 검토 모듈, 차트 필터·합계·키보드 선택, 모바일 표시, HTML 저장과 오프라인 복원을 검증했다. 기업 정확 일치는 52/135공고이며 미연결 83개를 유지한다.')
    p(d,'직무 후보 규칙의 의미 적합성 전수 검토, 최신 모집 상태, 기업 동일성 전체, 공식 방산 지정, 교육 품질과 비용 정의 확정은 후속 과제다. 사용자 탐색 개선이나 취업 성과를 실측한 결과로 현재 기대효과를 표현하지 않는다.')
    page(d,'11 이번 보완 이력')
    p(d,'원문 제목과 1~7절의 문장·표는 변경하지 않았다. 목적과 기대효과를 현재의 미구현 범위에 맞춰 축소하는 수정은 하지 않았다. 아래 네 절을 추가하여 기획과 구현 상태를 함께 읽을 수 있도록 했다.')
    table(d,['구분','위치','내용과 사유'],[
      ('유지','원문 1~7절','개요·배경·대상·기대효과·흐름·구성·활용 원칙 유지'),('추가','8절','채택 범위·시각화·검토·단일 선택·HTML 공유 현황'),('추가','9절','미채택 데이터와 미구현 비교 기능, 후보 규칙의 한계'),('추가','10절','7개 빈 영역·수집 우선순위·검증 완료 및 미검증 범위'),('추가','11절','변경 위치와 이유를 추적하는 보완 이력')],[2,3,12.7],10)
    p(d,'현재 기획안 자체의 핵심 방향을 바꿀 필요는 없다. 다만 상세 설계와 구현의 차이는 남아 있으므로 원문에 있던 기능이 구현되었다고 해석해서는 안 된다. 현재 버전의 선별 범위는 기획과 별도로 검토하고, 후속 추가 여부는 데이터 활용 목적과 기대효과에 따라 결정한다.')
    p(d,'함께 볼 문서: 01 데이터 명세와 활용설계, 02 제작과정과 기획정합성 보고서. 두 문서에는 파일별 핵심 필드·관측 단위·연결 규칙과 세부 차이·검증 결과가 기록되어 있다.')
    # The original doc's styles remain untouched; appended explicit runs use a locally available Korean font.
    new_nodes=[e for e in list(body) if e not in original and e.tag!=qn('w:sectPr')]
    from lxml import etree
    with zipfile.ZipFile(REF) as src:
        raw=src.read('word/document.xml');root=etree.fromstring(raw);target=root.find(qn('w:body'))
        sect=target.find(qn('w:sectPr'));at=list(target).index(sect) if sect is not None else len(target)
        for n in new_nodes:target.insert(at,deepcopy(n));at+=1
        out=OUT/'03_프로젝트_기획서_구현현황_보완본.docx'
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():dst.writestr(item,etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) if item.filename=='word/document.xml' else src.read(item.filename))
    with zipfile.ZipFile(REF) as a,zipfile.ZipFile(out) as b:
        changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
        assert changed==['word/document.xml'],changed
        aa=etree.fromstring(a.read('word/document.xml')).find(qn('w:body'))
        bb=etree.fromstring(b.read('word/document.xml')).find(qn('w:body'))
        old=[e for e in aa if e.tag!=qn('w:sectPr')]
        for x,y in zip(old,list(bb)):assert etree.tostring(x,method='c14n')==etree.tostring(y,method='c14n')
    (WORK/'plan-preservation.json').write_text(json.dumps({'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'changed_parts':changed,'original_body_elements_preserved':len(original_nonsect),'added_sections':['8','9','10','11']},ensure_ascii=False,indent=2),encoding='utf8')
    return out

if __name__=='__main__':
    for f in [build_spec(),build_report(),build_plan()]:print(f)
